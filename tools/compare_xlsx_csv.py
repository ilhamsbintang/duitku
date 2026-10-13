#!/usr/bin/env python3
"""Bandingkan workbook xlsx (Detail Transaksi) vs CSV yang dikirim user, per periode 10 Sep - 9 Okt 2026."""
import csv
import json
import sys
from collections import Counter
from datetime import date, datetime, timedelta

import openpyxl

XLSX = '/home/ubuntu/Laporan_Keuangan_Sep-Okt_2026.xlsx'
CSV_PATH = '/home/ubuntu/.hermes/cache/documents/doc_07d426d7e8f7_Laporan-Keuangan-2026-09 2.csv'


def xlsx_rows():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = None
    for name in wb.sheetnames:
        if 'Detail' in name:
            ws = wb[name]
            break
    if ws is None:
        sys.exit('sheet Detail tidak ketemu; sheets=' + str(wb.sheetnames))
    out = []
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        first = row[0]
        if isinstance(first, (datetime, date)):
            d = (first.date() if isinstance(first, datetime) else first).isoformat()
        elif isinstance(first, (int, float)):
            d = (date(1899, 12, 30) + timedelta(days=int(first))).isoformat()
        else:
            continue
        flow, desc, vendor, account, cat, sub, amount = row[2], row[3], row[4], row[5], row[6], row[7], row[8]
        note = ''
        if vendor and vendor != '-':
            note = vendor
        elif desc and desc != '-':
            note = desc
        out.append((d, 'in' if flow == 'CR' else 'out', abs(float(amount)), (cat or '').strip(),
                    (sub or '').strip() if sub and sub != '-' else '', (account or '').strip(), note))
    return out


def csv_rows():
    out = []
    with open(CSV_PATH, newline='', encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh, delimiter=';'):
            if not (r.get('Date') or '').strip():
                continue
            m, d, y = (int(x) for x in r['Date'].split('/'))
            vendor, desc = (r['Vendor'] or '').strip(), (r['Description'] or '').strip()
            vendor = '' if vendor == '-' else vendor
            desc = '' if desc == '-' else desc
            note = f'{vendor} — {desc}' if (vendor and desc) else (vendor or desc)
            out.append((f'{y:04d}-{m:02d}-{d:02d}', 'in' if r['Flow'].strip() == 'CR' else 'out',
                        abs(float(r['Amount'])), (r['Category'] or '').strip(),
                        (r['Sub-Category'] or '').strip(), (r['Financial account'] or '').strip(), note))
    return out


x, c = xlsx_rows(), csv_rows()
print(f'xlsx detail rows: {len(x)}   csv rows: {len(c)}')
print(f'xlsx in/out: {sum(r[2] for r in x if r[1]=="in"):,.2f} / {sum(r[2] for r in x if r[1]=="out"):,.2f}')
print(f'csv  in/out: {sum(r[2] for r in c if r[1]=="in"):,.2f} / {sum(r[2] for r in c if r[1]=="out"):,.2f}')
only_csv = Counter(c) - Counter(x)
only_xlsx = Counter(x) - Counter(c)
print('\nhanya ada di CSV  :')
for r, n in only_csv.items():
    print('  ', n, r)
print('hanya ada di XLSX :')
for r, n in only_xlsx.items():
    print('  ', n, r)
if not only_csv and not only_xlsx:
    print('\nIDENTIK — konversi 100% cocok dengan workbook.')
